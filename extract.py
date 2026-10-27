"""myTutor stage 2: turn every retrieved file in a job into clean Markdown text.

Usage
    py extract.py                     # the newest job in jobs_root
    py extract.py --day "Day 1"       # the newest job for Day 1 (add --course MBAM601 if needed)
    py extract.py --job 2026-10-17_192812_mbam601-...-day-1
    py extract.py --force             # extract again, even if text/ already exists

Reads   <job>/manifest.json and <job>/source/...
Writes  <job>/text/<same path>.md     one Markdown file per source, with page or slide headings
        <job>/extract.json            per file: pages, words, guessed kind, warnings, seconds

Needs:  py -m pip install -r requirements.txt   (pymupdf, python-pptx, python-docx)
"""

from __future__ import annotations

import argparse
import json
import os
import re
import sys
import time
from collections import Counter
from dataclasses import dataclass, field, asdict
from datetime import datetime, timezone
from pathlib import Path

LOW_TEXT_CHARS = 40           # a page with fewer characters is probably an image or a scan
REPEAT_SHARE = 0.6            # a line on this share of pages is a header/footer/watermark
SLIDE_WORDS_PER_PAGE = 120    # PDFs averaging fewer words per landscape page look like slides
FOLDER_KINDS = {"lecture": "slides", "lectures": "slides", "slides": "slides",
                "readings": "reading", "reading": "reading",
                "cases": "case", "case": "case"}


@dataclass
class Unit:
    """One page or one slide."""
    label: str                # "Page 3" or "Slide 3"
    text: str
    landscape: bool = False


@dataclass
class ExtractEntry:
    rel_path: str
    status: str               # extracted | skipped | failed
    text_path: str = ""
    kind_guess: str = ""      # slides | reading | case
    kind_reason: str = ""
    pages: int = 0
    words: int = 0
    low_text_pages: list = field(default_factory=list)
    removed_repeated_lines: int = 0
    seconds: float = 0.0
    note: str = ""


# ---------------------------------------------------------------- readers

def read_pdf(path: Path) -> list[Unit]:
    try:
        import pymupdf
    except ImportError:
        import fitz as pymupdf  # older PyMuPDF
    units = []
    with pymupdf.open(path) as doc:
        if doc.needs_pass:
            raise ValueError("PDF is password-protected")
        for i, page in enumerate(doc, start=1):
            rect = page.rect
            units.append(Unit(f"Page {i}", page.get_text("text", sort=True),
                              landscape=rect.width > rect.height))
    return units


def read_pptx(path: Path) -> list[Unit]:
    from pptx import Presentation
    prs = Presentation(str(path))
    units = []
    for i, slide in enumerate(prs.slides, start=1):
        lines = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                lines.extend(p.text for p in shape.text_frame.paragraphs if p.text.strip())
            elif getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    lines.append(" | ".join(c.text.strip() for c in row.cells))
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                lines.append(f"\nSpeaker notes: {notes}")
        units.append(Unit(f"Slide {i}", "\n".join(lines), landscape=True))
    return units


def read_docx(path: Path) -> list[Unit]:
    from docx import Document
    doc = Document(str(path))
    lines = []
    for p in doc.paragraphs:
        text = p.text.strip()
        if not text:
            continue
        style = (p.style.name or "").lower() if p.style is not None else ""
        if style.startswith("heading"):
            level = re.sub(r"\D", "", style) or "2"
            lines.append(f"{'#' * min(int(level) + 2, 6)} {text}")
        else:
            lines.append(text)
    for table in doc.tables:
        for row in table.rows:
            lines.append(" | ".join(c.text.strip() for c in row.cells))
    return [Unit("Document", "\n\n".join(lines))]


def read_plain(path: Path) -> list[Unit]:
    return [Unit("Document", path.read_text(encoding="utf-8", errors="replace"))]


READERS = {".pdf": read_pdf, ".pptx": read_pptx, ".docx": read_docx,
           ".txt": read_plain, ".md": read_plain}


# ---------------------------------------------------------------- cleaning

def clean_text(text: str) -> str:
    text = text.replace("­", "")                      # soft hyphens
    text = re.sub(r"(\w)-\n(\w)", r"\1\2", text)          # words split across lines
    text = re.sub(r"[ \t ]+", " ", text)
    lines = [ln.strip() for ln in text.splitlines()]
    text = "\n".join(lines)
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def remove_repeated_lines(units: list[Unit]) -> int:
    """Drop headers, footers and watermarks that repeat on most pages."""
    if len(units) < 4:
        return 0
    counts = Counter()
    for u in units:
        counts.update({ln for ln in u.text.splitlines() if ln.strip()})
    threshold = max(3, int(len(units) * REPEAT_SHARE))
    repeated = {ln for ln, n in counts.items() if n >= threshold and len(ln) > 3}
    # A line that is just a number is a page number: drop those too.
    removed = 0
    for u in units:
        kept = []
        for ln in u.text.splitlines():
            if ln in repeated or re.fullmatch(r"(page\s*)?\d{1,4}(\s*of\s*\d{1,4})?", ln.strip(), re.I):
                removed += 1
                continue
            kept.append(ln)
        u.text = "\n".join(kept).strip()
    return removed


# ---------------------------------------------------------------- kind guess

def guess_kind(rel_path: str, units: list[Unit], words: int) -> tuple[str, str]:
    parts = Path(rel_path).parts
    if len(parts) > 1 and parts[0].lower() in FOLDER_KINDS:
        return FOLDER_KINDS[parts[0].lower()], f"folder '{parts[0]}'"
    name = Path(rel_path).stem.lower()
    if Path(rel_path).suffix.lower() == ".pptx" or re.search(r"\b(slides?|lecture|session)\b", name):
        return "slides", "file name or type"
    if re.search(r"\bcase\b", name):
        return "case", "file name"
    if units:
        landscape = sum(u.landscape for u in units) / len(units)
        per_page = words / len(units)
        if landscape >= 0.6 and per_page < SLIDE_WORDS_PER_PAGE:
            return "slides", f"landscape pages, {per_page:.0f} words a page"
    return "reading", "default"


# ---------------------------------------------------------------- the run

def to_markdown(entry: ExtractEntry, origin: str, units: list[Unit]) -> str:
    title = Path(entry.rel_path).stem
    front = [
        "---",
        f"source: {json.dumps(entry.rel_path, ensure_ascii=False)}",
        f"origin: {json.dumps(origin, ensure_ascii=False)}",
        f"kind_guess: {entry.kind_guess}",
        f"pages: {entry.pages}",
        f"words: {entry.words}",
        "---",
        "",
        f"# {title}",
        "",
    ]
    body = []
    for u in units:
        if u.label != "Document":
            body.append(f"## {u.label}\n")
        body.append(u.text if u.text else "_(no text on this page)_")
        body.append("")
    return "\n".join(front + body).rstrip() + "\n"


def extract_job(job_dir: Path, force: bool = False) -> list[ExtractEntry]:
    manifest = json.loads((job_dir / "manifest.json").read_text(encoding="utf-8"))
    out_root = job_dir / "text"
    entries = []
    t_run = time.perf_counter()
    for f in manifest["files"]:
        rel = f["rel_path"]
        if f["status"] != "retrieved":
            continue
        src = job_dir / "source" / Path(rel)
        dest = out_root / Path(rel).with_suffix(".md")
        entry = ExtractEntry(rel_path=rel, status="")
        reader = READERS.get(src.suffix.lower())
        if reader is None:
            entry.status, entry.note = "skipped", f"no reader for {src.suffix}"
        elif dest.exists() and not force:
            entry.status, entry.note = "skipped", "already extracted (use --force to redo)"
            entry.text_path = dest.relative_to(job_dir).as_posix()
        else:
            t0 = time.perf_counter()
            try:
                units = reader(src)
                for u in units:
                    u.text = clean_text(u.text)
                entry.removed_repeated_lines = remove_repeated_lines(units)
                entry.pages = len(units)
                entry.words = sum(len(u.text.split()) for u in units)
                entry.low_text_pages = [i for i, u in enumerate(units, start=1)
                                        if len(u.text) < LOW_TEXT_CHARS]
                entry.kind_guess, entry.kind_reason = guess_kind(rel, units, entry.words)
                dest.parent.mkdir(parents=True, exist_ok=True)
                dest.write_text(to_markdown(entry, f.get("origin", ""), units), encoding="utf-8")
                entry.text_path = dest.relative_to(job_dir).as_posix()
                entry.status = "extracted"
                if entry.pages and len(entry.low_text_pages) == entry.pages:
                    entry.note = "no text found: probably a scan, needs OCR"
                elif entry.low_text_pages and entry.kind_guess != "slides":
                    entry.note = f"{len(entry.low_text_pages)} page(s) with little or no text"
            except Exception as exc:
                entry.status, entry.note = "failed", f"{type(exc).__name__}: {exc}"
            entry.seconds = round(time.perf_counter() - t0, 2)
        entries.append(entry)

    report = {
        "job": job_dir.name,
        "extracted_at": datetime.now(timezone.utc).astimezone().isoformat(timespec="seconds"),
        "seconds": round(time.perf_counter() - t_run, 2),
        "total_words": sum(e.words for e in entries if e.status == "extracted"),
        "files": [asdict(e) for e in entries],
    }
    tmp = job_dir / "extract.json.tmp"
    tmp.write_text(json.dumps(report, indent=2, ensure_ascii=False), encoding="utf-8")
    os.replace(tmp, job_dir / "extract.json")
    return entries


def latest_job(jobs_root: Path) -> Path:
    jobs = [p for p in jobs_root.iterdir() if (p / "manifest.json").exists()] if jobs_root.exists() else []
    if not jobs:
        sys.exit(f"No jobs in {jobs_root}: run retrieve.py first")
    return max(jobs, key=lambda p: (p / "manifest.json").stat().st_mtime)


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description="Extract text from a myTutor job.")
    parser.add_argument("--config", default="config.json")
    parser.add_argument("--job", help="job folder name (default: the newest)")
    parser.add_argument("--day", help='e.g. "Day 1": the newest job for that day')
    parser.add_argument("--course", help="e.g. MBAM601, when several courses have the same day")
    parser.add_argument("--force", action="store_true", help="extract again even if text exists")
    args = parser.parse_args(argv)

    cfg_path = Path(args.config)
    cfg = json.loads(cfg_path.read_text(encoding="utf-8")) if cfg_path.exists() else {}
    jobs_root = Path(os.path.expandvars(os.path.expanduser(cfg.get("jobs_root", "~/mytutor/jobs"))))
    from jobs import find_job
    job_dir = find_job(jobs_root, args.job, args.day, args.course, needs="manifest.json")
    if not (job_dir / "manifest.json").exists():
        sys.exit(f"Not a job folder: {job_dir}")

    entries = extract_job(job_dir, force=args.force)
    for e in entries:
        detail = f"{e.pages:>3} pp {e.words:>6} words  {e.kind_guess:<7}" if e.status == "extracted" else ""
        note = f"  ({e.note})" if e.note else ""
        print(f"{e.status:<10} {detail} {e.rel_path}{note}")
    done = [e for e in entries if e.status == "extracted"]
    print(f"\n{len(done)} extracted, {sum(e.words for e in done):,} words, "
          f"{sum(e.status == 'failed' for e in entries)} failed")
    print(f"Text folder: {job_dir / 'text'}")
    return 1 if any(e.status == "failed" for e in entries) else 0


if __name__ == "__main__":
    sys.exit(main())
