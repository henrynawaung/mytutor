"""Tests for extract.py. Run from the mytutor folder:  py -m unittest discover tests"""

import json
import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import extract  # noqa: E402

try:
    import pymupdf  # noqa: F401
    HAVE_PDF = True
except ImportError:
    try:
        import fitz  # noqa: F401
        HAVE_PDF = True
    except ImportError:
        HAVE_PDF = False

try:
    import pptx  # noqa: F401
    import docx  # noqa: F401
    HAVE_OFFICE = True
except ImportError:
    HAVE_OFFICE = False


def make_pdf(path: Path, pages: list, landscape=False, marks=True):
    try:
        import pymupdf
    except ImportError:
        import fitz as pymupdf
    doc = pymupdf.open()
    w, h = (842, 595) if landscape else (595, 842)
    for i, text in enumerate(pages, start=1):
        page = doc.new_page(width=w, height=h)
        if marks:
            page.insert_text((50, 40), "For the exclusive use of H. Student, 2026.", fontsize=8)
            page.insert_text((w / 2, h - 30), str(i), fontsize=8)
        if text:
            page.insert_textbox(pymupdf.Rect(50, 80, w - 50, h - 80), text, fontsize=11)
    doc.save(path)


def make_job(root: Path, files: dict) -> Path:
    job = root / "job1"
    (job / "source").mkdir(parents=True)
    entries = []
    for rel, writer in files.items():
        dest = job / "source" / rel
        dest.parent.mkdir(parents=True, exist_ok=True)
        writer(dest)
        entries.append({"rel_path": rel, "status": "retrieved", "origin": f"C:/OneDrive/{rel}"})
    entries.append({"rel_path": "extra.zip", "status": "skipped_type", "origin": ""})
    (job / "manifest.json").write_text(json.dumps({"files": entries}))
    return job


class CleaningTests(unittest.TestCase):
    def test_clean_joins_hyphenated_words_and_spaces(self):
        self.assertEqual(extract.clean_text("trans-\nformation   is\u00a0 here\n\n\n\nnext"),
                         "transformation is here\n\nnext")

    def test_repeated_header_and_page_numbers_removed(self):
        units = [extract.Unit(f"Page {i}", f"HBR header\nBody {i}\n{i}") for i in range(1, 6)]
        removed = extract.remove_repeated_lines(units)
        self.assertEqual(units[2].text, "Body 3")
        self.assertEqual(removed, 10)

    def test_short_documents_keep_repeated_lines(self):
        units = [extract.Unit("Page 1", "Title\nA"), extract.Unit("Page 2", "Title\nB")]
        self.assertEqual(extract.remove_repeated_lines(units), 0)

    def test_kind_guess(self):
        reading = [extract.Unit("Page 1", "", landscape=False)] * 10
        slides = [extract.Unit("Page 1", "", landscape=True)] * 10
        self.assertEqual(extract.guess_kind("Cases/Kodak.pdf", reading, 5000)[0], "case")
        self.assertEqual(extract.guess_kind("Session 1 Slides Final.pdf", reading, 5000)[0], "slides")
        self.assertEqual(extract.guess_kind("Short Case.pdf", reading, 900)[0], "case")
        self.assertEqual(extract.guess_kind("deck.pdf", slides, 600)[0], "slides")
        self.assertEqual(extract.guess_kind("The New Disrupters.pdf", reading, 6000)[0], "reading")


@unittest.skipUnless(HAVE_PDF and HAVE_OFFICE, "needs pymupdf, python-pptx, python-docx")
class ExtractJobTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        root = Path(self.tmp.name)
        bodies = [f"Section {n}: digital transformation changes how firms create value, "
                  f"argument {n} of the article. " * 6 for n in range(1, 6)]

        def pptx_writer(p):
            from pptx import Presentation
            prs = Presentation()
            s = prs.slides.add_slide(prs.slide_layouts[1])
            s.shapes.title.text = "Session 1"
            s.placeholders[1].text = "Why digital transformation fails"
            s.notes_slide.notes_text_frame.text = "Ask the class for examples"
            prs.save(p)

        def docx_writer(p):
            from docx import Document
            d = Document()
            d.add_heading("Case: Kodak", level=1)
            d.add_paragraph("Kodak invented the digital camera in 1975.")
            d.save(p)

        self.job = make_job(root, {
            "The New Disrupters.pdf": lambda p: make_pdf(p, bodies),
            "deck.pdf": lambda p: make_pdf(p, ["Agenda", "Three waves", "Q&A", "Close"], landscape=True),
            "scan.pdf": lambda p: make_pdf(p, [""] * 2, marks=False),
            "Lecture/Session 1.pptx": pptx_writer,
            "Cases/Kodak.docx": docx_writer,
            "broken.pdf": lambda p: p.write_bytes(b"not a pdf"),
        })

    def tearDown(self):
        self.tmp.cleanup()

    def entry(self, entries, rel):
        return next(e for e in entries if e.rel_path == rel)

    def test_extracts_each_type(self):
        entries = extract.extract_job(self.job)
        reading = self.entry(entries, "The New Disrupters.pdf")
        self.assertEqual((reading.status, reading.pages, reading.kind_guess), ("extracted", 5, "reading"))
        md = (self.job / reading.text_path).read_text(encoding="utf-8")
        self.assertIn("## Page 3", md)
        self.assertIn('origin: "C:/OneDrive/The New Disrupters.pdf"', md)
        self.assertNotIn("exclusive use", md)          # watermark removed

        slides = self.entry(entries, "Lecture/Session 1.pptx")
        md = (self.job / slides.text_path).read_text(encoding="utf-8")
        self.assertIn("Speaker notes: Ask the class", md)
        self.assertEqual(slides.kind_guess, "slides")

        case = self.entry(entries, "Cases/Kodak.docx")
        self.assertEqual(case.kind_guess, "case")
        self.assertIn("### Case: Kodak", (self.job / case.text_path).read_text(encoding="utf-8"))

        self.assertEqual(self.entry(entries, "deck.pdf").kind_guess, "slides")

    def test_scans_and_broken_files_are_reported(self):
        entries = extract.extract_job(self.job)
        self.assertIn("OCR", self.entry(entries, "scan.pdf").note)
        broken = self.entry(entries, "broken.pdf")
        self.assertEqual(broken.status, "failed")
        self.assertNotIn("extra.zip", [e.rel_path for e in entries])

    def test_report_and_rerun(self):
        extract.extract_job(self.job)
        report = json.loads((self.job / "extract.json").read_text(encoding="utf-8"))
        self.assertGreater(report["total_words"], 300)
        again = extract.extract_job(self.job)
        self.assertEqual(self.entry(again, "deck.pdf").status, "skipped")
        forced = extract.extract_job(self.job, force=True)
        self.assertEqual(self.entry(forced, "deck.pdf").status, "extracted")


if __name__ == "__main__":
    unittest.main()
